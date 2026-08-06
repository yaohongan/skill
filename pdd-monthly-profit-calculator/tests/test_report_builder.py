import json
import os
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


SKILL_DIR = Path(__file__).resolve().parents[1]


def find_node():
    configured = os.environ.get("CODEX_NODE")
    if configured:
        return Path(configured)
    system_node = shutil.which("node")
    if system_node:
        return Path(system_node)
    candidates = list(Path.home().glob(
        ".cache/codex-runtimes/*/dependencies/node/bin/node.exe"
    ))
    return candidates[0] if candidates else None


def find_node_modules():
    configured = os.environ.get("CODEX_NODE_MODULES")
    if configured:
        return Path(configured)
    candidates = list(Path.home().glob(
        ".cache/codex-runtimes/*/dependencies/node/node_modules"
    ))
    return candidates[0] if candidates else None


class ReportBuilderTests(unittest.TestCase):
    def test_builds_traceable_chinese_workbook_with_editable_ad_spend(self):
        node = find_node()
        node_modules = find_node_modules()
        if not node or not node.exists() or not node_modules or not node_modules.exists():
            self.skipTest("未找到 Codex Node.js 或工作区 Node 模块")
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            data_path = root / "分析.json"
            output_path = root / "客户核对表.xlsx"
            data_path.write_text(json.dumps(self.sample_data(), ensure_ascii=False), encoding="utf-8")
            env = os.environ.copy()
            env["CODEX_NODE_MODULES"] = str(node_modules)
            subprocess.run(
                [str(node), str(SKILL_DIR / "scripts" / "build_report.mjs"),
                 "--input", str(data_path), "--output", str(output_path)],
                check=True,
                env=env,
                capture_output=True,
                text=True,
                encoding="utf-8",
            )

            self.assertTrue(output_path.exists())
            workbook = load_workbook(output_path, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                ["核算说明", "商品盈亏汇总", "有效订单明细", "成本映射", "未匹配商品", "排除订单"],
            )
            overview = workbook["核算说明"]
            self.assertEqual(overview["B10"].value, 12)
            self.assertEqual(overview["B13"].value, "=B9-B10")
            self.assertEqual(overview["B10"].fill.fgColor.rgb[-6:], "FFF1B8")
            self.assertIn("暂算", overview["A1"].value)
            self.assertIn("｜利润核算", overview["A1"].value)
            self.assertEqual(workbook["有效订单明细"]["L3"].value, "=H3-K3")
            self.assertEqual(workbook["有效订单明细"]["F3"].value, 1001)
            self.assertEqual(workbook["有效订单明细"]["F3"].number_format, "0")

    @staticmethod
    def sample_data():
        detail = {
            "商品": "测试商品", "订单号": "A", "订单状态": "已收货", "售后状态": "无售后或售后取消",
            "订单成交时间": "2026-07-10 12:00:00", "商品id": "1001", "商品规格": "默认",
            "商家实收金额(元)": 30, "商品数量(件)": 2, "单件总成本": 10,
            "订单总成本": 20, "商品毛利": 10,
        }
        return {
            "shop_name": "测试店铺", "start_date": "2026-07-01", "end_date": "2026-07-30",
            "cost_path": "成本.xlsx", "order_path": "订单.csv", "ad_spend_input": 12,
            "ad_spend_is_approximate": True, "valid_rows": 2, "matched_valid_rows": 1, "valid_qty": 2,
            "excluded_status_rows": 0, "merchant_received": 30, "product_cost": 20,
            "gross_profit": 10, "net_profit_after_ads": -2, "unmatched_rows": 1,
            "unmatched_qty": 1, "unmatched_received": 25, "is_temporary": True,
            "cost_duplicate_same_ids": [], "sku_summary": [{"商品id": "1001", "有效订单行数": 1,
                "有效数量": 2, "商家实收": 30, "单件总成本": 10, "订单总成本": 20,
                "商品毛利": 10, "毛利率": 0.3333}],
            "matched_order_details": [detail], "excluded_order_details": [],
            "unmatched_order_details": [{key: value for key, value in detail.items() if key not in {"单件总成本", "订单总成本", "商品毛利"}}],
            "unmatched_products": [{"商品id": "9999", "订单行数": 1, "商品数量": 1, "商家实收": 25}],
            "cost_source_rows": [{"商品id": "1001", "总成本": 10}],
        }


if __name__ == "__main__":
    unittest.main()

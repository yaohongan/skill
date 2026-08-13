import importlib.util
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


SKILL_ROOT = Path(__file__).resolve().parents[1]
BUILDER = SKILL_ROOT / "scripts" / "build_profit_report.py"


def load_builder():
    spec = importlib.util.spec_from_file_location("build_profit_report", BUILDER)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def sample_data():
    return {
        "状态": "完成",
        "参数": {
            "店铺名称": "测试店铺", "统计开始日期": "2026-07-01", "统计结束日期": "2026-07-30",
            "时间字段": "支付时间", "时间口径说明": "优先使用支付时间", "推广总花费": 30,
            "推广费来源": "测试截图", "推广截图": "截图.png", "订单文件": "订单.csv",
            "订单数据来源": "CSV", "成本文件": "成本.xlsx",
        },
        "数据概况": {"区间订单数": 1, "有效订单数": 1, "退款取消数": 0, "售后处理中数": 0, "剔除订单数": 0, "唯一规格映射数": 1, "未匹配规格数": 0},
        "正式汇总": {"商家实收": 100, "商品总成本": 40, "推广前利润": 60, "推广总花费": 30, "最终盈亏": 30, "净利润率": 0.3},
        "订单明细": [{"序号": 1, "订单号": "A1", "商品ID": "123456789012345678", "订单时间": "2026-07-02 10:00:00", "订单状态": "已完成", "售后状态": "", "商品名称": "南瓜子", "商品规格": "原味350g", "商品数量": 2, "商家实收": 100, "成本工作表": "成本", "成本SKU名称": "原味350g", "单件总成本": 20, "匹配方式": "标准化名称精确匹配"}],
        "SKU映射": [{"商品ID": "123456789012345678", "商品名称": "南瓜子", "商品规格": "原味350g", "有效订单数": 1, "商品件数": 2, "商家实收": 100, "成本工作表": "成本", "成本SKU名称": "原味350g", "单件总成本": 20, "匹配状态": "已匹配", "匹配方式": "标准化名称精确匹配"}],
        "成本标准": [{"工作表": "成本", "行号": 2, "口味说明": "原味", "成本SKU名称": "原味350g", "重量克数": 350, "分类": "南瓜子", "口味": "原味", "成本": 17, "运费": 3, "总成本": 20}],
        "剔除订单": [],
        "未匹配SKU": [],
    }


class ReportContractTests(unittest.TestCase):
    def test_builder_contains_all_required_chinese_sheets(self):
        module = load_builder()
        self.assertEqual(module.SHEET_NAMES, ["客户汇报", "SKU汇总", "订单明细", "成本标准", "剔除订单", "未匹配SKU", "参数与来源", "核对检查"])

    def test_builder_shows_ad_spend_as_a_separate_deduction(self):
        source = BUILDER.read_text(encoding="utf-8")
        self.assertIn("减：推广总花费", source)
        self.assertIn("推广前利润-推广总花费", source)

    def test_builder_creates_reopenable_workbook_with_formulas(self):
        module = load_builder()
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "客户版.xlsx"
            module.build_workbook(sample_data(), output)
            workbook = load_workbook(output, data_only=False)
            self.assertEqual(workbook.sheetnames, module.SHEET_NAMES)
            self.assertEqual(workbook["客户汇报"]["B14"].value, "='参数与来源'!B8")
            self.assertEqual(workbook["客户汇报"]["B15"].value, "=B13-B14")
            self.assertEqual(workbook["订单明细"]["C2"].value, "123456789012345678")
            self.assertEqual(workbook["订单明细"]["N2"].value, "=I2*M2")
            workbook.close()


if __name__ == "__main__":
    unittest.main()

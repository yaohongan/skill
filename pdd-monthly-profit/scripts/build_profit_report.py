import argparse
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SHEET_NAMES = [
    "客户汇报",
    "SKU汇总",
    "订单明细",
    "成本标准",
    "剔除订单",
    "未匹配SKU",
    "参数与来源",
    "核对检查",
]

NAVY = "17324D"
BLUE = "2563EB"
GREEN = "138A72"
RED = "C53B3B"
AMBER = "D28A00"
WHITE = "FFFFFF"
PALE_BLUE = "EAF2FF"
PALE_GREEN = "EAF7F3"
PALE_RED = "FDEBEC"
PALE_AMBER = "FFF5DB"
LINE = "D7DEE8"
MONEY = '#,##0.00;[Red](#,##0.00);-'
COUNT = '#,##0;[Red](#,##0);-'
THIN = Side(style="thin", color=LINE)


def parse_datetime(value):
    if not value:
        return None
    return datetime.fromisoformat(str(value))


def set_title(sheet, end_column, text):
    sheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=end_column)
    cell = sheet.cell(1, 1, text)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name="Microsoft YaHei", size=16, bold=True, color=WHITE)
    cell.alignment = Alignment(vertical="center")
    for row in sheet.iter_rows(min_row=1, max_row=2, min_col=1, max_col=end_column):
        for item in row:
            item.fill = PatternFill("solid", fgColor=NAVY)


def add_table(sheet, headers, rows, start_row=1):
    for column, value in enumerate(headers, 1):
        cell = sheet.cell(start_row, column, value)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_index, values in enumerate(rows, start_row + 1):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_index, column, value)
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    sheet.freeze_panes = sheet.cell(start_row + 1, 1)
    sheet.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{max(start_row, start_row + len(rows))}"


def set_widths(sheet, widths):
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width


def preserve_id(cell, value):
    cell.value = str(value or "")
    cell.number_format = "@"


def build_parameters(workbook, data):
    sheet = workbook["参数与来源"]
    sheet.sheet_view.showGridLines = False
    set_title(sheet, 4, "核算参数与数据来源")
    params = data["参数"]
    overview = data["数据概况"]
    totals = data["正式汇总"]
    rows = [
        ["店铺名称", params["店铺名称"], "用户提供/文件识别", "用于报告标题"],
        ["统计开始日期", parse_datetime(params["统计开始日期"]), "投流截图", "含当日"],
        ["统计结束日期", parse_datetime(params["统计结束日期"]), "投流截图", "含当日"],
        ["订单时间字段", params["时间字段"], "订单文件", params["时间口径说明"]],
        ["推广总花费(元)", params["推广总花费"], params.get("推广费来源", "拼多多投流数据截图"), "只使用总花费，不使用成交花费"],
        ["推广截图", params.get("推广截图") or "未记录路径", "用户提供", "截图用于确认日期和总花费"],
        ["订单文件", params["订单文件"], params["订单数据来源"], "月度订单源文件"],
        ["成本文件", params["成本文件"], "成本Excel全部工作表", "使用SKU名称和总成本"],
        ["区间订单数", overview["区间订单数"], "订单文件", "含剔除订单"],
        ["有效订单数", overview["有效订单数"], "订单文件", "已排除退款、取消及售后处理中"],
        ["退款/取消订单数", overview["退款取消数"], "订单文件", "不计收入和成本"],
        ["售后处理中订单数", overview["售后处理中数"], "订单文件", "不计收入和成本"],
        ["剔除订单数", overview["剔除订单数"], "订单文件", "退款/取消加售后处理中"],
        ["唯一规格映射数", overview["唯一规格映射数"], "有效订单", "按商品ID和商品规格统计"],
        ["未匹配规格数", overview["未匹配规格数"], "成本匹配", "正式结果必须为0"],
        ["源数据商家实收(元)", totals["商家实收"], "计算结果JSON", "用于Excel勾稽"],
        ["源数据商品总成本(元)", totals["商品总成本"], "计算结果JSON", "用于Excel勾稽"],
        ["源数据推广前利润(元)", totals["推广前利润"], "计算结果JSON", "用于Excel勾稽"],
        ["源数据最终盈亏(元)", totals["最终盈亏"], "计算结果JSON", "用于Excel勾稽"],
    ]
    add_table(sheet, ["参数", "数值", "来源", "说明"], rows, 3)
    sheet["B5"].number_format = "yyyy-mm-dd"
    sheet["B6"].number_format = "yyyy-mm-dd"
    for row in range(8, 9):
        sheet.cell(row, 2).number_format = MONEY
    for row in range(19, 23):
        sheet.cell(row, 2).number_format = MONEY
    sheet["B8"].comment = Comment("推广费用采用截图中的‘总花费（元）’，不采用‘成交花费（元）’。", "Skill")
    set_widths(sheet, [27, 47, 35, 54])


def build_order_detail(workbook, data):
    sheet = workbook["订单明细"]
    headers = ["序号", "订单号", "商品ID", "订单时间", "订单状态", "售后状态", "商品名称", "商品规格", "商品数量", "商家实收金额(元)", "成本工作表", "匹配成本SKU", "单件总成本(元)", "订单总成本(元)", "推广前利润(元)", "匹配方式"]
    rows = []
    for row in data["订单明细"]:
        rows.append([row["序号"], row["订单号"], str(row["商品ID"]), parse_datetime(row["订单时间"]), row["订单状态"], row["售后状态"], row["商品名称"], row["商品规格"], row["商品数量"], row["商家实收"], row["成本工作表"], row["成本SKU名称"], row["单件总成本"], None, None, row["匹配方式"]])
    add_table(sheet, headers, rows)
    for index, row in enumerate(data["订单明细"], 2):
        preserve_id(sheet.cell(index, 3), row["商品ID"])
        sheet.cell(index, 14, f"=I{index}*M{index}")
        sheet.cell(index, 15, f"=J{index}-N{index}")
        sheet.cell(index, 4).number_format = "yyyy-mm-dd hh:mm:ss"
        for column in (10, 13, 14, 15):
            sheet.cell(index, column).number_format = MONEY
    set_widths(sheet, [9, 24, 24, 22, 19, 19, 42, 42, 14, 18, 24, 35, 18, 18, 18, 28])


def build_sku_summary(workbook, data):
    sheet = workbook["SKU汇总"]
    headers = ["商品ID", "商品名称", "商品规格", "有效订单数", "商品件数", "商家实收金额(元)", "成本工作表", "匹配成本SKU", "单件总成本(元)", "商品总成本(元)", "推广前利润(元)", "匹配状态", "匹配方式"]
    rows = [[str(r["商品ID"]), r["商品名称"], r["商品规格"], r["有效订单数"], r["商品件数"], r["商家实收"], r["成本工作表"], r["成本SKU名称"], r["单件总成本"], None, None, r["匹配状态"], r["匹配方式"]] for r in data["SKU映射"]]
    add_table(sheet, headers, rows)
    for index, row in enumerate(data["SKU映射"], 2):
        preserve_id(sheet.cell(index, 1), row["商品ID"])
        sheet.cell(index, 10, f"=E{index}*I{index}")
        sheet.cell(index, 11, f"=F{index}-J{index}")
        for column in (6, 9, 10, 11):
            sheet.cell(index, column).number_format = MONEY
    set_widths(sheet, [24, 42, 42, 16, 16, 18, 24, 35, 18, 18, 18, 16, 28])


def build_costs(workbook, data):
    sheet = workbook["成本标准"]
    headers = ["工作表", "行号", "口味说明", "成本SKU名称", "重量(g)", "分类", "口味", "商品成本(元)", "运费(元)", "总成本(元)"]
    rows = [[r["工作表"], r["行号"], r["口味说明"], r.get("成本SKU名称") or "（成本表未填写SKU名称）", r["重量克数"], r["分类"], r["口味"], r["成本"], r["运费"], r["总成本"]] for r in data["成本标准"]]
    add_table(sheet, headers, rows)
    for row in range(2, len(rows) + 2):
        for column in (8, 9, 10):
            sheet.cell(row, column).number_format = MONEY
    set_widths(sheet, [24, 10, 38, 42, 14, 18, 18, 18, 18, 18])


def build_excluded(workbook, data):
    sheet = workbook["剔除订单"]
    headers = ["序号", "订单号", "商品ID", "订单时间", "订单状态", "售后状态", "商品名称", "商品规格", "商品数量", "商家实收金额(元)", "剔除原因"]
    rows = [[r["序号"], r["订单号"], str(r["商品ID"]), parse_datetime(r["订单时间"]), r["订单状态"], r["售后状态"], r["商品名称"], r["商品规格"], r["商品数量"], r["商家实收"], r["剔除原因"]] for r in data["剔除订单"]]
    add_table(sheet, headers, rows)
    for index, row in enumerate(data["剔除订单"], 2):
        preserve_id(sheet.cell(index, 3), row["商品ID"])
        sheet.cell(index, 4).number_format = "yyyy-mm-dd hh:mm:ss"
        sheet.cell(index, 10).number_format = MONEY
    set_widths(sheet, [9, 24, 24, 22, 20, 20, 42, 42, 14, 18, 18])


def build_unmatched(workbook, data):
    sheet = workbook["未匹配SKU"]
    sheet.sheet_view.showGridLines = False
    set_title(sheet, 8, "未匹配SKU清单")
    headers = ["商品ID", "商品名称", "商品规格", "有效订单数", "商品件数", "商家实收金额(元)", "未匹配原因", "处理状态"]
    source = data["未匹配SKU"]
    rows = [[str(r["商品ID"]), r["商品名称"], r["商品规格"], r["有效订单数"], r["商品件数"], r["商家实收"], r["未匹配原因"], "等待补成本"] for r in source]
    if not rows:
        rows = [["", "", "无", 0, 0, 0, "全部SKU已匹配", "已完成"]]
    add_table(sheet, headers, rows, 4)
    for index, row in enumerate(source, 5):
        preserve_id(sheet.cell(index, 1), row["商品ID"])
    set_widths(sheet, [24, 42, 42, 16, 16, 18, 36, 18])


def build_summary(workbook, data):
    sheet = workbook["客户汇报"]
    sheet.sheet_view.showGridLines = False
    params = data["参数"]
    totals = data["正式汇总"]
    month = params["统计开始日期"][:7].replace("-", "年") + "月"
    set_title(sheet, 8, f"{month}{params['店铺名称']}盈亏汇报")
    sheet.merge_cells("A4:H4")
    sheet["A4"] = f"核算期间：{params['统计开始日期']} 至 {params['统计结束日期']}    有效订单：{data['数据概况']['有效订单数']:,}笔"
    sheet["A4"].fill = PatternFill("solid", fgColor=PALE_BLUE)
    sheet["A4"].font = Font(name="Microsoft YaHei", size=11, bold=True, color=BLUE)
    sheet.merge_cells("A6:B8")
    sheet["A6"] = "最终结果"
    sheet.merge_cells("C6:E8")
    sheet["C6"] = "=B15"
    sheet["C6"].number_format = MONEY
    sheet.merge_cells("F6:H8")
    sheet["F6"] = '=IF(B15>=0,"盈利","亏损")'
    result_fill = PALE_GREEN if totals["最终盈亏"] >= 0 else PALE_RED
    result_color = GREEN if totals["最终盈亏"] >= 0 else RED
    for cell in (sheet["A6"], sheet["C6"]):
        cell.fill = PatternFill("solid", fgColor=result_fill)
        cell.font = Font(name="Microsoft YaHei", size=16, bold=True, color=result_color)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet["F6"].fill = PatternFill("solid", fgColor=PALE_AMBER)
    sheet["F6"].font = Font(name="Microsoft YaHei", size=16, bold=True, color=AMBER)
    sheet["F6"].alignment = Alignment(horizontal="center", vertical="center")
    labels = ["商家实收金额", "减：商品总成本", "推广前利润", "减：推广总花费", "最终盈亏", "净利润率", "盈亏状态"]
    notes = ["有效订单实际到账", "单件总成本×商品数量", "尚未扣除推广费", "截图中的总花费（元）", "推广前利润-推广总花费", "最终盈亏÷商家实收", "根据最终盈亏自动判断"]
    formulas = ["=SUM('订单明细'!J2:J1048576)", "=SUM('订单明细'!N2:N1048576)", "=B11-B12", "='参数与来源'!B8", "=B13-B14", "=IFERROR(B15/B11,0)", '=IF(B15>=0,"盈利","亏损")']
    add_table(sheet, ["核算项目", "金额/比例", "说明"], list(zip(labels, formulas, notes)), 10)
    for index, formula in enumerate(formulas, 11):
        sheet.cell(index, 2, formula)
    for row in range(11, 16):
        sheet.cell(row, 2).number_format = MONEY
    sheet["B16"].number_format = "0.0%"
    sheet["B14"].comment = Comment("推广总花费已在此处单独扣除，并进入最终盈亏公式。", "Skill")
    set_widths(sheet, [31, 24, 48, 17, 17, 17, 17, 17])


def build_checks(workbook, data):
    sheet = workbook["核对检查"]
    sheet.sheet_view.showGridLines = False
    set_title(sheet, 5, "核对检查")
    overview = data["数据概况"]
    totals = data["正式汇总"]
    detail_received = round(sum(float(r["商家实收"]) for r in data["订单明细"]), 2)
    detail_cost = round(sum(float(r["商品数量"]) * float(r["单件总成本"]) for r in data["订单明细"]), 2)
    checks = [
        ["区间订单数=有效订单+剔除订单", overview["区间订单数"], overview["有效订单数"] + overview["剔除订单数"]],
        ["有效订单数=订单明细行数", overview["有效订单数"], len(data["订单明细"])],
        ["订单明细实收=源数据实收", detail_received, totals["商家实收"]],
        ["订单明细成本=源数据成本", detail_cost, totals["商品总成本"]],
        ["推广前利润=源数据推广前利润", round(detail_received - detail_cost, 2), totals["推广前利润"]],
        ["最终盈亏=推广前利润-推广费", round(totals["推广前利润"] - totals["推广总花费"], 2), totals["最终盈亏"]],
        ["未匹配SKU数=0", overview["未匹配规格数"], 0],
    ]
    rows = []
    for label, calculated, target in checks:
        difference = round(float(calculated) - float(target), 8)
        rows.append([label, calculated, target, difference, "PASS" if abs(difference) < 0.01 else "FAIL"])
    add_table(sheet, ["检查项目", "计算值", "目标值", "差异", "状态"], rows, 3)
    sheet["A12"] = "模型状态"
    sheet["B12"] = "PASS" if all(row[4] == "PASS" for row in rows) else "FAIL"
    set_widths(sheet, [48, 22, 22, 22, 15])


def validate_workbook(path):
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        if workbook.sheetnames != SHEET_NAMES:
            raise ValueError("生成的工作表不完整或顺序错误")
        if workbook["核对检查"]["B12"].value != "PASS":
            raise ValueError("核对检查未通过")
        if workbook["客户汇报"]["B14"].value != "='参数与来源'!B8":
            raise ValueError("推广费扣除公式缺失")
    finally:
        workbook.close()


def build_workbook(data, output_path):
    if data.get("状态") != "完成" or not data.get("正式汇总"):
        raise ValueError("存在未匹配SKU或其他阻断项，不能生成正式客户版Excel")
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in SHEET_NAMES:
        workbook.create_sheet(name)
    build_summary(workbook, data)
    build_sku_summary(workbook, data)
    build_order_detail(workbook, data)
    build_costs(workbook, data)
    build_excluded(workbook, data)
    build_unmatched(workbook, data)
    build_parameters(workbook, data)
    build_checks(workbook, data)
    workbook.save(output)
    validate_workbook(output)
    return output.resolve()


def build_parser():
    parser = argparse.ArgumentParser(description="生成拼多多月度盈亏中文客户版Excel")
    parser.add_argument("--data", required=True, help="prepare_profit_data.py生成的JSON路径")
    parser.add_argument("--output", required=True, help="输出XLSX路径")
    return parser


def main():
    args = build_parser().parse_args()
    data = json.loads(Path(args.data).read_text(encoding="utf-8"))
    output = build_workbook(data, args.output)
    print(json.dumps({"状态": "完成", "输出Excel": str(output), "工作表": SHEET_NAMES}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
